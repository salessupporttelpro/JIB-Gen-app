"""
build_docx.py
Build a JIB (Justifikasi Inisiatif Bisnis) .docx from structured JSON data,
following the standard Telkom Property table layout (No | Aspek | Isi).
"""

import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

GLOBAL_FONT = "Calibri"

ROW_LABELS = [
    ("1", "Latar Belakang Kebutuhan Bisnis", "latar_belakang"),
    (
        "2",
        "Ruang Lingkup Pekerjaan, termasuk Jumlah Kebutuhan, persyaratan, lokasi deployment dan rencana pelaksanaan",
        "ruang_lingkup",
    ),
    ("3", "Asumsi yang digunakan", "asumsi"),
    ("4", "Analisis-Analisis dan Mitigasi Risiko", "analisis"),
    ("5", "Kesimpulan", "kesimpulan"),
    ("6", "Rekomendasi Keputusan", "rekomendasi_keputusan"),
]


def _find_local_logo() -> str:
    """Mencari file gambar logo berdasarkan lokasi absolut build_docx.py"""
    # Mengambil lokasi folder tempat script build_docx.py berada
    base_dir = os.path.dirname(os.path.abspath(__file__))

    possible_paths = [
        os.path.join(base_dir, "references", "logo.png"),
        os.path.join(base_dir, "references", "logo.jpg"),
        os.path.join(base_dir, "references", "logo.PNG"),
        os.path.join(base_dir, "logo.png"),
    ]

    for path in possible_paths:
        if os.path.exists(path):
            return path
    return None


def _clean_text(text: str, add_dot: True) -> str:
    """Menghapus penomoran ganda di awal dan opsional menambahkan titik di akhir."""
    if not text:
        return "-"
    s = str(text).strip()

    # Hapus prefix angka/bullet di awal (misal "1. ", "a) ", "- ")
    cleaned = re.sub(r"^(?:(?:\d+|[a-zA-Z])[\.\)]\s*|[-•]\s*)+", "", s).strip()
    if not cleaned:
        cleaned = s

    # Tambahkan titik di akhir kalimat HANYA JIKA add_dot=True (untuk narasi)
    if add_dot and cleaned and cleaned[-1] not in [".", "!", "?", ":", ";"]:
        cleaned += "."

    return cleaned


def _set_cell_shading(cell, color_hex="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _clear_cell(cell):
    cell.text = ""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    return p


def _add_line_direct(
    cell, first_para, text, bold=False, level=0, size=10, new_para=True
):
    """Menambahkan teks langsung ke sel Word tanpa memicu pembersihan regex."""
    if new_para:
        p = cell.add_paragraph()
    else:
        p = first_para
    p.paragraph_format.left_indent = Cm(0.4 * level)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run(str(text))
    run.font.name = GLOBAL_FONT
    run.font.size = Pt(size)
    run.bold = bold
    return p


def _set_table_borders(table, color_hex="000000", sz="4", val="single"):
    """Menambahkan garis border pada tabel anak secara tegas."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")

    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), val)
        border.set(qn("w:sz"), sz)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color_hex)
        tblBorders.append(border)

    tblPr.append(tblBorders)


def _add_flexible_table(cell, rows):
    """Membuat tabel anak (tabel finansial/kelayakan) dengan ukuran kolom memadai."""
    if not isinstance(rows, list) or not rows:
        return

    sample = rows[0]
    if isinstance(sample, dict):
        keys = list(sample.keys())
        keys = [k for k in keys if k not in ("bold", "level")]
    else:
        keys = ["label", "value"]

    num_cols = len(keys) if keys else 2

    tbl = cell.add_table(rows=len(rows) + 1, cols=num_cols)
    tbl.autofit = False

    # PERBAIKAN LEBAR KOLOM: Beri ruang lebih luas untuk nominal angka
    if num_cols == 2:
        col_widths = [Inches(2.6), Inches(1.8)]  # Kolom nilai dilebarkan dari 0.75" ke 1.8"
    else:
        col_widths = [Inches(4.4) / num_cols] * num_cols

    for i, col in enumerate(tbl.columns):
        if i < len(col_widths):
            col.width = col_widths[i]

    _set_table_borders(tbl, color_hex="000000", sz="4")

    # Header Tabel
    headers = ["Uraian Pekerjaan", "Remarks"] if num_cols == 2 else ["Item", "Nilai"]
    hdr_cells = tbl.rows[0].cells
    for j, h_text in enumerate(headers[:num_cols]):
        hdr_cells[j].text = h_text
        _set_cell_shading(hdr_cells[j], "D9D9D9")
        for p in hdr_cells[j].paragraphs:
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            for r in p.runs:
                r.font.name = GLOBAL_FONT
                r.font.size = Pt(9)
                r.bold = True

    # Isi Data
    for i, item in enumerate(rows, start=1):
        is_bold = isinstance(item, dict) and item.get("bold", False)
        for j, k in enumerate(keys):
            c = tbl.cell(i, j)
            if j < len(col_widths):
                c.width = col_widths[j]

            if isinstance(item, dict):
                val_text = str(item.get(k, "-"))
            else:
                val_text = str(item)
            
            # PERBAIKAN TITIK: Set add_dot=False agar isi sel tabel tidak diberi titik di akhir
            c.text = _clean_text(val_text, add_dot=False)

            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.name = GLOBAL_FONT
                    r.font.size = Pt(9)
                    if is_bold:
                        r.bold = True


def _add_page_number_fields(run):
    """Menambahkan field dinamis nomor halaman pada Word XML."""
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def _fill_list_cell(cell, items, default_numbering=True):
    first = _clear_cell(cell)
    first_used = False
    if not items:
        _add_line_direct(cell, first, "-", new_para=False)
        return

    main_idx = 0
    sub_idx = 0

    for item in items:
        if isinstance(item, dict):
            text = str(item.get("text", "-"))
            level = item.get("level", 0)
        else:
            text = str(item) if item else "-"
            level = 0

        clean_content = _clean_text(text)

        if default_numbering:
            if level == 0:
                main_idx += 1
                sub_idx = 0
                full_text = f"{main_idx}. {clean_content}"
            else:
                prefix_char = chr(97 + (sub_idx % 26))
                full_text = f"{prefix_char}. {clean_content}"
                sub_idx += 1
        else:
            full_text = clean_content

        _add_line_direct(cell, first, full_text, level=level, new_para=first_used)
        first_used = True


def _fill_text_cell(cell, text):
    first = _clear_cell(cell)
    text_str = str(text).strip() if text else "-"
    for i, line in enumerate(text_str.split("\n")):
        _add_line_direct(cell, first, line, new_para=(i > 0))


def build_jib_docx(data: dict, output_path: str):
    doc = Document()

    # 0. Global Document Font Settings
    style = doc.styles["Normal"]
    style.font.name = GLOBAL_FONT
    style.font.size = Pt(10)

    # 1. Header Dokumen
    logo_path = _find_local_logo()

    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False

    cell_logo = header_table.cell(0, 0)
    cell_info = header_table.cell(0, 1)
    cell_logo.width = Inches(1.5)
    cell_info.width = Inches(5.0)

    if logo_path:
        p_logo = cell_logo.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(1.4))

    p_info = cell_info.paragraphs[0]
    p_info.paragraph_format.space_after = Pt(2)

    r1 = p_info.add_run("PT Graha Sarana Duta\n")
    r1.font.name = GLOBAL_FONT
    r1.bold = True
    r1.font.size = Pt(12)

    r2 = p_info.add_run("Jl. Kebon Sirih no. 10 Jakarta\n\n")
    r2.font.name = GLOBAL_FONT
    r2.bold = True
    r2.font.size = Pt(12)

    r3 = p_info.add_run("JUSTIFIKASI INISIATIF BISNIS/ ")
    r3.font.name = GLOBAL_FONT
    r3.bold = True
    r3.font.size = Pt(12)

    r4 = p_info.add_run("SPECIAL BUSINESS REQUEST\n")
    r4.font.name = GLOBAL_FONT
    r4.bold = True
    r4.font.size = Pt(12)
    r4.font.strike = True

    judul_p = str(data.get("judul_project", "-")).strip()
    r5 = p_info.add_run(f"Project : {judul_p}")
    r5.font.name = GLOBAL_FONT
    r5.bold = True
    r5.font.size = Pt(12)

    doc.add_paragraph()

    # Metadata UNIT INISIATOR & NO. DRP/DRK
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)

    r_unit_lbl = p_meta.add_run("UNIT INISIATOR\t: ")
    r_unit_lbl.font.name = GLOBAL_FONT
    r_unit_lbl.bold = True

    r_unit_val = p_meta.add_run(f"{str(data.get('unit_kerja_inisiator', '-'))}\n")
    r_unit_val.font.name = GLOBAL_FONT

    r_drp_lbl = p_meta.add_run("NO. DRP/DRK\t: ")
    r_drp_lbl.font.name = GLOBAL_FONT
    r_drp_lbl.bold = True

    # 2. Tabel Utama JIB (6 Baris x 3 Kolom)
    table = doc.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(0.4), Inches(1.6), Inches(4.5)]

    for no, label, key in ROW_LABELS:
        row = table.add_row()
        for cell, w in zip(row.cells, widths):
            cell.width = w
        row.cells[0].text = no
        row.cells[1].text = label

        for cell_idx in [0, 1]:
            for p_ in row.cells[cell_idx].paragraphs:
                for r_ in p_.runs:
                    r_.font.name = GLOBAL_FONT
                    if cell_idx == 1:
                        r_.bold = True

        content_cell = row.cells[2]

        if key == "latar_belakang":
            _fill_list_cell(
                content_cell, data.get("latar_belakang", []), default_numbering=True
            )

        elif key == "ruang_lingkup":
            _fill_list_cell(
                content_cell, data.get("ruang_lingkup", []), default_numbering=True
            )

        elif key == "asumsi":
            asumsi = data.get("asumsi", {})
            first = _clear_cell(content_cell)

            fin_table = (
                asumsi.get("financial_table", []) if isinstance(asumsi, dict) else []
            )
            if fin_table and isinstance(fin_table, list) and len(fin_table) > 0:
                _add_line_direct(
                    content_cell,
                    first,
                    "Perhitungan nilai pekerjaan",
                    bold=True,
                    new_para=False,
                )
                _add_flexible_table(content_cell, fin_table)

            catatan_list = asumsi.get("catatan", []) if isinstance(asumsi, dict) else []
            if catatan_list:
                for i, line in enumerate(catatan_list):
                    clean_line = _clean_text(str(line))
                    prefix_char = chr(97 + (i % 26))
                    full_line = f"{prefix_char}. {clean_line}"
                    _add_line_direct(content_cell, first, full_line, new_para=True)
            elif not fin_table:
                _add_line_direct(content_cell, first, "-", new_para=False)

        elif key == "analisis":
            analisis = (
                data.get("analisis", {})
                if isinstance(data.get("analisis"), dict)
                else {}
            )
            first = _clear_cell(content_cell)
            first_used = False

            data_kelayakan = (
                analisis.get("kelayakan_proyek")
                or analisis.get("kelayakan")
                or analisis.get("aspek_kelayakan")
                or []
            )

            sections = [
                ("1. ASPEK STRATEGIS", analisis.get("strategis", [])),
                ("2. ASPEK KOMPETISI", analisis.get("kompetisi", [])),
                ("3. ASPEK MANFAAT", analisis.get("manfaat", [])),
                ("4. ASPEK BISNIS", analisis.get("bisnis", [])),
                ("5. ASPEK KELAYAKAN PROYEK", data_kelayakan),
                ("6. ASPEK TEKNIS", analisis.get("teknis", [])),
            ]

            for section_title, items in sections:
                _add_line_direct(
                    content_cell,
                    first,
                    section_title,
                    bold=True,
                    new_para=first_used,
                )
                first_used = True

                if isinstance(items, str):
                    items = [items]

                if not items:
                    if section_title != "5. ASPEK KELAYAKAN PROYEK":
                        _add_line_direct(
                            content_cell, first, "-", level=1, new_para=True
                        )
                else:
                    for idx, it in enumerate(items):
                        clean_item = _clean_text(str(it))
                        prefix_char = chr(97 + (idx % 26))
                        full_str = f"{prefix_char}. {clean_item}"
                        _add_line_direct(
                            content_cell,
                            first,
                            full_str,
                            level=1,
                            new_para=True,
                        )

                if section_title == "5. ASPEK KELAYAKAN PROYEK":
                    kt = analisis.get("kelayakan_table", [])
                    if kt and isinstance(kt, list) and len(kt) > 0:
                        _add_flexible_table(content_cell, kt)

            # 7. ANALISIS DAN MITIGASI RISIKO
            risko_mitigasi_list = analisis.get("risiko_dan_mitigasi", [])
            if risko_mitigasi_list and isinstance(risko_mitigasi_list, list):
                _add_line_direct(content_cell, first, "", new_para=True)
                _add_line_direct(
                    content_cell,
                    first,
                    "7. ANALISIS DAN MITIGASI RISIKO",
                    bold=True,
                    new_para=True,
                )

                for item_rm in risko_mitigasi_list:
                    if isinstance(item_rm, dict):
                        kat_title = item_rm.get("kategori", "Risiko")
                        _add_line_direct(
                            content_cell,
                            first,
                            f"{kat_title}:",
                            bold=True,
                            level=1,
                            new_para=True,
                        )

                        r_list = item_rm.get("risiko", [])
                        if r_list:
                            _add_line_direct(
                                content_cell,
                                first,
                                "Risiko:",
                                bold=True,
                                level=2,
                                new_para=True,
                            )
                            for i_r, r_text in enumerate(r_list):
                                prefix_char = chr(97 + (i_r % 26))
                                _add_line_direct(
                                    content_cell,
                                    first,
                                    f"{prefix_char}. {_clean_text(str(r_text))}",
                                    level=3,
                                    new_para=True,
                                )

                        m_list = item_rm.get("mitigasi", [])
                        if m_list:
                            _add_line_direct(
                                content_cell,
                                first,
                                "Mitigasi:",
                                bold=True,
                                level=2,
                                new_para=True,
                            )
                            for i_m, m_text in enumerate(m_list):
                                prefix_char = chr(97 + (i_m % 26))
                                _add_line_direct(
                                    content_cell,
                                    first,
                                    f"{prefix_char}. {_clean_text(str(m_text))}",
                                    level=3,
                                    new_para=True,
                                )

        elif key == "kesimpulan":
            kesimpulan_obj = data.get("kesimpulan", {})
            first = _clear_cell(content_cell)

            if isinstance(kesimpulan_obj, dict):
                status_ai = _clean_text(
                    str(kesimpulan_obj.get("status_kelayakan", "[Layak/Tidak Layak]"))
                )
                judul_p = str(data.get("judul_project", "-")).strip()

                # Jika AI tidak menyertakan kata Layak/Tidak Layak, gunakan placeholder [Layak/Tidak Layak]
                if "layak" not in status_ai.lower():
                    status_ai = f"[Layak/Tidak Layak] {status_ai}".strip()

                pembuka = f"Kesimpulan: Project {judul_p} dinyatakan {status_ai} dengan catatan untuk memperhatikan risiko-risiko sebagai berikut :"
                _add_line_direct(content_cell, first, pembuka, new_para=False)

                poin_list = kesimpulan_obj.get("poin_kesimpulan", [])
                if isinstance(poin_list, list):
                    for i, p_text in enumerate(poin_list, 1):
                        clean_p = _clean_text(str(p_text))
                        _add_line_direct(
                            content_cell,
                            first,
                            f"{i}.  {clean_p}",
                            level=1,
                            new_para=True,
                        )
            else:
                _add_line_direct(
                    content_cell, first, str(kesimpulan_obj), new_para=False
                )

        elif key == "rekomendasi_keputusan":
            _fill_text_cell(content_cell, data.get("rekomendasi_keputusan", "-"))

        elif key == "keputusan":
            _fill_text_cell(content_cell, data.get("keputusan", "-"))

    doc.add_paragraph()

    # 3. Tabel Tanda Tangan
    roles_list = [
        ("Inisiator", data.get("inisiator")),
        ("Reviewer", data.get("reviewer")),
        ("Approver", data.get("approver")),
    ]

    sig_table = doc.add_table(rows=1 + len(roles_list), cols=4)
    sig_table.style = "Table Grid"

    headers = ["KET", "NAMA", "JABATAN", "TANDA TANGAN"]
    for idx_header, text_header in enumerate(headers):
        cell_hdr = sig_table.rows[0].cells[idx_header]
        cell_hdr.text = text_header
        _set_cell_shading(cell_hdr, "D9D9D9")
        for para_hdr in cell_hdr.paragraphs:
            for run_hdr in para_hdr.runs:
                run_hdr.font.name = GLOBAL_FONT
                run_hdr.bold = True

    for row_idx, (role_name, person_info) in enumerate(roles_list, start=1):
        cells_row = sig_table.rows[row_idx].cells
        cells_row[0].text = role_name

        if isinstance(person_info, dict):
            cells_row[1].text = str(person_info.get("nama", "-") or "-")
            cells_row[2].text = str(person_info.get("jabatan", "-") or "-")
        elif isinstance(person_info, str):
            cells_row[1].text = person_info if person_info else "-"
            cells_row[2].text = "-"
        else:
            cells_row[1].text = "-"
            cells_row[2].text = "-"

        cells_row[3].text = ""

        for col_idx in range(4):
            for p_sig in cells_row[col_idx].paragraphs:
                for r_sig in p_sig.runs:
                    r_sig.font.name = GLOBAL_FONT

    # 4. Footer Otomatis (Calibri)
    section = doc.sections[0]
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.text = ""  # Bersihkan teks default

    # Set paragraf footer langsung Rata Kanan (Right Alignment)
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # 1. Judul Project
    r_ftr_title = p_ftr.add_run(f"{judul_p} | ")
    r_ftr_title.font.name = GLOBAL_FONT
    r_ftr_title.font.size = Pt(8.5)
    r_ftr_title.font.italic = False

    # 2. Angka Nomor Halaman
    r_ftr_page = p_ftr.add_run()
    r_ftr_page.font.name = GLOBAL_FONT
    r_ftr_page.font.size = Pt(8.5)
    _add_page_number_fields(r_ftr_page)

    doc.save(output_path)
    return output_path

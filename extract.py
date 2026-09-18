"""
extract.py -- Modul ekstraksi teks dan tabel dari berbagai format file (PDF, XLSX, DOCX)
baik dari unggahan pengguna Streamlit maupun folder referensi lokal.
"""

import io
import os
import docx
import openpyxl
import pdfplumber


def extract_pdf_text(file_bytes: bytes, max_chars: int = 15000) -> str:
    """Mengekstraksi seluruh teks dan tabel dari file PDF dengan batasan jumlah karakter."""
    content_parts = []
    total_len = 0

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            # Stop iterasi halaman jika batas karakter sudah tercapai
            if max_chars and total_len >= max_chars:
                break

            # 1. Ekstrak teks halaman
            text = page.extract_text()
            if text:
                content_parts.append(text)
                total_len += len(text)

            # 2. Ekstrak tabel halaman (jika ada)
            tables = page.extract_tables()
            for table in tables:
                if max_chars and total_len >= max_chars:
                    break
                for row in table:
                    clean_row = [
                        str(cell).strip() if cell is not None else "" for cell in row
                    ]
                    if any(clean_row):
                        row_str = " | ".join(clean_row)
                        content_parts.append(row_str)
                        total_len += len(row_str)

    full_text = "\n".join(content_parts)

    # Potong string secara presisi jika melebihi max_chars
    if max_chars and len(full_text) > max_chars:
        return full_text[:max_chars] + "\n... (data dibatasi/truncated)"

    return full_text


def extract_xlsx_text(file_bytes: bytes, max_rows: int = 500) -> str:
    """Mengekstraksi dan meratakan lembar kerja Excel (RAB) menjadi format teks terstruktur berpemisah '|'."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []

    for ws in wb.worksheets:
        lines.append(f"--- Sheet: {ws.title} ---")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            cells = [str(c).strip() if c is not None else "" for c in row]
            lines.append(" | ".join(cells))
            row_count += 1
            if row_count >= max_rows:
                lines.append("... (data dibatasi/truncated)")
                break

    return "\n".join(lines)


def extract_reference_files(uploaded_files: dict) -> str:
    """
    Membaca dan menggabungkan seluruh file transaksi unggahan pengguna Streamlit (PDF & XLSX)
    menjadi satu blok teks utuh sebagai sumber data transaksi utama.
    """
    sections = []
    for name, data in uploaded_files.items():
        lower_name = name.lower()
        try:
            if lower_name.endswith(".pdf"):
                content = extract_pdf_text(data)
            elif lower_name.endswith((".xlsx", ".xlsm")):
                content = extract_xlsx_text(data)
            else:
                content = "(tipe file tidak didukung, dilewati)"
        except Exception as e:
            content = f"(gagal membaca file: {e})"

        sections.append(f"===== FILE TRANSAKSI: {name} =====\n{content}\n")

    return "\n".join(sections)


def read_all_local_references(folder_path: str = "references") -> str:
    """
    Memindai dan membaca seluruh file PDF & DOCX dari folder referensi lokal
    sebagai acuan struktur dan gaya penulisan korporat.
    """
    if not os.path.exists(folder_path):
        return ""

    content_list = []

    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)

        if not os.path.isfile(file_path):
            continue

        ext = os.path.splitext(filename)[1].lower()

        # 1. Ekstrak dari PDF lokal
        if ext == ".pdf":
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        text = page.extract_text()
                        if text:
                            content_list.append(text)

                        tables = page.extract_tables()
                        for table in tables:
                            for row in table:
                                clean_row = [
                                    str(cell).strip() if cell is not None else ""
                                    for cell in row
                                ]
                                if any(clean_row):
                                    content_list.append(" | ".join(clean_row))
            except Exception as e:
                print(f"Error membaca PDF {filename}: {e}")

        # 2. Ekstrak dari Word (.docx) lokal
        elif ext == ".docx":
            try:
                doc = docx.Document(file_path)

                # Ekstrak Paragraf
                for para in doc.paragraphs:
                    if para.text.strip():
                        content_list.append(para.text.strip())

                # Ekstrak Tabel Word
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        ]
                        if row_text:
                            content_list.append(" | ".join(row_text))
            except Exception as e:
                print(f"Error membaca DOCX {filename}: {e}")

    return "\n".join(content_list)
